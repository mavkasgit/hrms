"""ОДНОРАЗОВЫЙ скрипт восстановления (инцидент 2026-08-27/28).

Контекст: маунты data/notifications и data/statements отсутствовали в
docker-compose.prod.yml, файлы .docx жили в эфемерном слое контейнера и
погибли при пересоздании 2026-08-27 15:47. БД цела (строки + file_path),
шаблоны целы (data/templates примонтирован).

Скрипт воспроизводит пайплайн create_draft (document_draft_service):
тип/сотрудник -> замены -> шаблон -> render_docx_placeholders -> docx,
но пишет файл ТОЧНО по file_path из БД (без _unique_file_path).
Файлы не перезаписывает (EXISTS -> skip). Запуск:
    docker cp scripts/rescue_regen_docx.py hrms-backend-prod:/tmp/rescue_regen.py
    docker exec -e PYTHONPATH=/app hrms-backend-prod python /tmp/rescue_regen.py
"""

import asyncio

from sqlalchemy import select
from sqlalchemy.orm import joinedload

from app.core.database import async_session
from app.core.paths import storage_root
from app.models.employee import Employee
from app.models.notification import Notification
from app.models.statement import Statement
from app.services.document_draft_service import (
    notification_draft_service,
    statement_draft_service,
)
from app.services.docx_renderer import load_template_or_create_blank, render_docx_placeholders


async def regen_kind(label: str, svc, model, type_attr: str) -> None:
    cfg = svc._config()

    def blank_doc():
        from docx import Document

        return Document()

    async with async_session() as db:
        rows = (await db.execute(select(model).order_by(model.id))).scalars().all()
        print(f"== {label}: {len(rows)} rows ==")
        for row in rows:
            if not row.file_path:
                print(f"  id={row.id} SKIP (file_path is NULL)")
                continue
            target = cfg.path_func(row.file_path)
            if target.exists():
                print(f"  id={row.id} EXISTS {target.name}")
                continue

            type_obj = None
            type_id = getattr(row, type_attr, None)
            if type_id is not None:
                type_obj = await cfg.type_getter(db, type_id)

            employee = None
            if row.employee_id is not None:
                res = await db.execute(
                    select(Employee)
                    .options(joinedload(Employee.position), joinedload(Employee.department))
                    .where(Employee.id == row.employee_id)
                )
                employee = res.scalar_one_or_none()

            replacements = cfg.replacements_builder(
                title=row.title,
                number=row.number or "",
                doc_date=row.date,
                employee=employee,
                type_obj=type_obj,
                extra_fields=row.extra_fields,
            )

            # Семантика _load_document: template_filename типа -> его шаблон
            # (нет файла -> пустой документ); иначе default template.docx
            # вида; иначе пустой документ.
            template_filename = getattr(type_obj, "template_filename", None) if type_obj else None
            if type_obj and template_filename:
                template_path = cfg.template_getter(type_obj)
                doc = (
                    await load_template_or_create_blank(template_path)
                    if template_path.exists()
                    else blank_doc()
                )
            else:
                default_template = storage_root(cfg.dir_key) / "template.docx"
                doc = (
                    await load_template_or_create_blank(default_template)
                    if default_template.exists()
                    else blank_doc()
                )

            render_docx_placeholders(doc, replacements)
            target.parent.mkdir(parents=True, exist_ok=True)
            await asyncio.to_thread(doc.save, str(target))
            print(f"  id={row.id} REGEN {target.name} (template={template_filename or 'default/blank'})")


async def main() -> None:
    await regen_kind("notifications", notification_draft_service, Notification, "notification_type_id")
    await regen_kind("statements", statement_draft_service, Statement, "statement_type_id")
    print("DONE")


if __name__ == "__main__":
    asyncio.run(main())
